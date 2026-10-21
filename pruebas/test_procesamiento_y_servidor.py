"""Pruebas del procesamiento del Excel real y del servidor (de punta a punta por HTTP)."""
import glob
import hashlib
import json
import os
import shutil
import sys
import tempfile
import threading
import unittest
import urllib.error
import urllib.request
from http.server import ThreadingHTTPServer

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(RAIZ, "app"))
import procesar_excel  # noqa: E402
import servidor  # noqa: E402

# El Excel original tiene datos personales y no está en el repositorio público: si falta,
# solo se omiten las pruebas que lo leen (las del juego y del servidor usan los datos procesados).
EXCELES = [p for p in glob.glob(os.path.join(RAIZ, "*.xlsx")) if not os.path.basename(p).startswith("~$")]
EXCEL = EXCELES[0] if EXCELES else None
JSON_PROCESADO = os.path.join(RAIZ, "datos", "preguntas_procesadas.json")


def sha(ruta):
    with open(ruta, "rb") as f:
        return hashlib.sha256(f.read()).hexdigest()


@unittest.skipUnless(EXCEL, "Falta el Excel original en la carpeta principal (no se publica en GitHub por privacidad)")
class PruebasProcesamiento(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.hash_antes = sha(EXCEL)
        cls.res = procesar_excel.procesar(EXCEL)
        cls.excel = procesar_excel.leer_excel(EXCEL)

    def test_excel_original_no_se_modifica(self):
        self.assertEqual(sha(EXCEL), self.hash_antes)

    def test_estructura(self):
        self.assertEqual(self.res["personas_que_respondieron"], 41)
        self.assertEqual(len(self.res["preguntas"]), 51)
        self.assertEqual([o["numero"] for o in self.res["preguntas_omitidas"]], [29])
        self.assertEqual(self.res["avisos"], [])

    def test_cada_respuesta_original_se_cuenta_exactamente_una_vez(self):
        _, preguntas_excel, _ = self.excel
        por_num = {p["numero"]: p for p in preguntas_excel}
        for p in self.res["preguntas"]:
            filas = [o["fila"] for a in p["respuestas"] for o in a["originales"]] + [e["fila"] for e in p["excluidas"]]
            esperadas = [r["fila"] for r in por_num[p["numero"]]["respuestas"] if r["texto"] and r["texto"].strip()]
            self.assertEqual(sorted(filas), sorted(esperadas), f"Pregunta {p['numero']}")
            self.assertEqual(p["respuestas_validas"] + p["respuestas_excluidas"], len(esperadas))
            textos = {r["fila"]: r["texto"].strip() for r in por_num[p["numero"]]["respuestas"] if r["texto"]}
            privados = procesar_excel.leer_nombres_privados().get(str(p["numero"]), {})
            for a in p["respuestas"]:
                for o in a["originales"]:  # trazabilidad exacta (con los nombres de personas ocultos)
                    self.assertEqual(o["texto"], procesar_excel.anonimizar(textos[o["fila"]], privados.get(a["etiqueta"])))

    def test_frecuencias_reales_y_orden(self):
        for p in self.res["preguntas"]:
            frecs = [a["frecuencia"] for a in p["respuestas"]]
            self.assertEqual(frecs, sorted(frecs, reverse=True))
            for a in p["respuestas"]:
                self.assertEqual(a["frecuencia"], len(a["originales"]))
            self.assertEqual(sum(frecs), p["respuestas_validas"])

    def test_archivo_procesado_al_dia(self):
        with open(JSON_PROCESADO, encoding="utf-8") as f:
            guardado = json.load(f)
        quitar = lambda d: [{k: v for k, v in p.items()} for p in d["preguntas"]]  # noqa: E731
        self.assertEqual(quitar(guardado), quitar(self.res))

    def test_ejemplos_conocidos(self):
        p = {x["numero"]: x for x in self.res["preguntas"]}
        self.assertEqual(p[4]["respuestas"][0]["etiqueta"], "Había trancón")
        self.assertEqual(p[4]["respuestas"][0]["frecuencia"], 28)
        tablero4 = [a for a in p[4]["respuestas"] if a["en_tablero"]]
        self.assertEqual([a["frecuencia"] for a in tablero4], [28, 4, 2, 2])   # P4 mejorada: 4 respuestas
        self.assertEqual([a["puntos"] for a in tablero4], procesar_excel.repartir_100([28, 4, 2, 2]))
        self.assertEqual(p[6]["respuestas_excluidas"], 10)  # respuestas fuera de tema en «tipo de oso»

    # ---------------------------------------------------------------- puntuación: 100 por pregunta
    def test_reparto_de_100_casos_conocidos(self):
        r = procesar_excel.repartir_100
        self.assertEqual(r([30, 25, 20, 15, 7, 3]), [30, 25, 20, 15, 7, 3])   # el ejemplo de la solicitud
        self.assertEqual(r([5]), [100])                                       # una sola respuesta
        self.assertEqual(r([3, 3, 3, 3]), [25, 25, 25, 25])                   # empate perfecto
        self.assertEqual(sum(r([2, 2, 2])), 100)                              # 33,3 cada una: suma 100 igual
        self.assertEqual(r([28, 2, 2]), [88, 6, 6])
        self.assertEqual(r([10, 6, 4, 2]), [46, 27, 18, 9])
        self.assertEqual(r([150, 50]), [75, 25])                              # frecuencias grandes: proporción
        with self.assertRaises(ValueError):
            r([0, 3])

    def test_reparto_de_100_en_miles_de_casos_aleatorios(self):
        import random
        azar = random.Random(2026)
        for _ in range(20000):
            k = azar.randint(1, 8)
            fr = sorted((azar.randint(1, azar.choice([5, 60, 500])) for _ in range(k)), reverse=True)
            pts = procesar_excel.repartir_100(fr)
            self.assertEqual(sum(pts), 100, fr)
            self.assertTrue(all(isinstance(x, int) and x >= 1 for x in pts), (fr, pts))
            for i in range(k - 1):
                if fr[i] > fr[i + 1]:
                    self.assertGreaterEqual(pts[i], pts[i + 1], (fr, pts))

    def test_cada_pregunta_habilitada_suma_exactamente_100(self):
        for p in self.res["preguntas"]:
            tablero = [a for a in p["respuestas"] if a["en_tablero"]]
            if not p["habilitada"]:
                continue
            self.assertEqual(sum(a["puntos"] for a in tablero), 100, p["numero"])
            self.assertEqual(p["puntos_tablero"], 100)
            self.assertTrue(4 <= len(tablero) <= 10)
            for a in tablero:
                self.assertIsInstance(a["puntos"], int)
                self.assertGreaterEqual(a["puntos"], 1)
            for a in p["respuestas"]:
                if not a["en_tablero"]:
                    self.assertEqual(a["puntos"], 0)          # fuera del tablero: no puntúa
                    self.assertTrue(a["motivo_fuera"])

    def test_tableros_con_distinta_cantidad_de_respuestas_suman_100(self):
        por_tamano = {}
        for p in self.res["preguntas"]:
            n = p["respuestas_en_tablero"]
            por_tamano.setdefault(n, []).append(sum(a["puntos"] for a in p["respuestas"] if a["en_tablero"]))
        for n in (4, 5, 6, 8, 10):                             # existen tableros de 4, 5, 6, 8 y 10 respuestas
            self.assertIn(n, por_tamano)
            self.assertEqual(set(por_tamano[n]), {100}, n)

    def test_la_mas_popular_tiene_mas_puntos_y_el_orden_se_conserva(self):
        empates_partidos = []
        for p in self.res["preguntas"]:
            rs = [a for a in p["respuestas"] if a["en_tablero"]]
            self.assertEqual(rs[0]["puntos"], max(a["puntos"] for a in rs), p["numero"])
            for x, y in zip(rs, rs[1:]):
                self.assertGreaterEqual(x["puntos"], y["puntos"], p["numero"])
                if x["frecuencia"] > y["frecuencia"]:
                    self.assertGreater(x["puntos"], y["puntos"], p["numero"])
                elif x["puntos"] != y["puntos"]:
                    empates_partidos.append(p["numero"])
        # Único caso inevitable: P1 tiene un triple empate en el primer lugar y 100 no se divide en partes iguales.
        self.assertEqual(sorted(set(empates_partidos)), [1])

    def test_excel_consolidado_tiene_los_mismos_puntos_que_usa_el_juego(self):
        import openpyxl
        wb = openpyxl.load_workbook(os.path.join(RAIZ, "datos", "preguntas_procesadas.xlsx"), read_only=True)
        filas = list(wb["Consolidadas"].iter_rows(values_only=True))
        enc = filas[0]
        i_num, i_pos = enc.index("N°"), enc.index("Posición")
        i_freq, i_pts = enc.index("Frecuencia real (personas)"), enc.index("PUNTOS ASIGNADOS (suman 100 por pregunta)")
        i_etq, i_tab = enc.index("Respuesta consolidada (en pantalla)"), enc.index("¿Va al tablero?")
        with open(JSON_PROCESADO, encoding="utf-8") as f:
            juego = {p["numero"]: p for p in json.load(f)["preguntas"]}
        sumas = {}
        for fila in filas[1:]:
            a = juego[fila[i_num]]["respuestas"][fila[i_pos] - 1]
            self.assertEqual((fila[i_etq], fila[i_freq], fila[i_pts]), (a["etiqueta"], a["frecuencia"], a["puntos"]))
            self.assertEqual(fila[i_tab] == "Sí", a["en_tablero"])
            sumas[fila[i_num]] = sumas.get(fila[i_num], 0) + fila[i_pts]
        self.assertEqual(set(sumas.values()), {100})              # en el Excel también suma 100 cada pregunta
        self.assertEqual(len(filas) - 1, sum(len(p["respuestas"]) for p in juego.values()))
        hoja_preg = list(wb["Preguntas"].iter_rows(values_only=True))
        col = hoja_preg[0].index("SUMA DE PUNTOS DEL TABLERO")
        self.assertEqual({f[col] for f in hoja_preg[1:]}, {100})
        self.assertIn("Contenido", wb.sheetnames)

    # ---------------------------------------------------------------- humor y contenido
    def buscar(self, numero, etiqueta):
        p = next(x for x in self.res["preguntas"] if x["numero"] == numero)
        return p, next((a for a in p["respuestas"] if a["etiqueta"] == etiqueta), None)

    def test_el_humor_cotidiano_y_las_groserias_se_conservan(self):
        for numero, etiqueta in ((35, "Comérselo"), (46, "Comerse los mocos"), (1, "Sacarse los mocos"), (17, "Tirarse un pedo"),
                                 (21, "¡Hijueputa!"), (27, "Mierda"), (52, "Pornografía"), (12, "Juguetes sexuales"),
                                 (13, "Un consolador"), (2, "Vaya coma mierda")):
            _, a = self.buscar(numero, etiqueta)
            self.assertIsNotNone(a, etiqueta)
            self.assertIsNone(a["revision_contenido"], etiqueta)
        _, a = self.buscar(21, "¡Hijueputa!")
        self.assertEqual(a["frecuencia"], 23)  # las variantes de la grosería se consolidaron sin perder frecuencia
        palabras = ("pedo", "moco", "mierda", "porno", "hp", "hijueputa", "consolador", "dildo", "gonorrea", "chimba")
        for p in self.res["preguntas"]:  # ninguna exclusión se debe a la grosería en sí
            for e in p["excluidas"]:
                self.assertFalse(any(w in e["texto"].lower() for w in palabras), e)

    def test_contenido_fuerte_se_marca_para_revision_sin_perder_frecuencia(self):
        p, a = self.buscar(15, "Una mujer")
        self.assertIsNotNone(a["revision_contenido"])
        self.assertEqual(p["respuestas_validas"], 41)  # sigue contando en la frecuencia
        for numero, etiqueta in ((5, "A los venezolanos"), (12, "Un venezolano"), (5, "A los de otra religión")):
            _, a = self.buscar(numero, etiqueta)
            self.assertTrue(a["revision_contenido"], etiqueta)
        import juego as motor
        for p in self.res["preguntas"]:  # nunca aparecen en el tablero mientras estén en revisión
            ids = {a["id"] for a in motor.Juego.tablero_de(p)}
            for a in p["respuestas"]:
                if a["revision_contenido"]:
                    self.assertNotIn(a["id"], ids)

    def test_reformulaciones_registradas_y_originales_intactos(self):
        _, a = self.buscar(23, "¡Hágame un catorce!")
        self.assertEqual(a["reformulacion"]["etiqueta_original"], "[nombre], un catorce")
        self.assertEqual([o["texto"] for o in a["originales"]], ["[nombre] un catorce"])
        # Los nombres reales están solo en el archivo privado (no publicado): no pueden aparecer en nada generado.
        privados = procesar_excel.leer_nombres_privados()
        nombres = [n for q in privados.values() for lista in q.values() for n in lista]
        self.assertTrue(nombres, "falta datos/privado/nombres_a_ocultar.json")
        generado = json.dumps(self.res, ensure_ascii=False)
        for n in nombres:
            self.assertNotRegex(generado, r"(?i)\b" + n + r"\b")


class PruebasServidor(unittest.TestCase):
    """Levanta el servidor real en un puerto libre y lo usa como lo haría el navegador."""

    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.partida = os.path.join(self.tmp, "partida.json")
        self.servicio = servidor.Servicio(JSON_PROCESADO, self.partida)
        self.httpd = ThreadingHTTPServer(("127.0.0.1", 0), servidor.crear_manejador(self.servicio))
        self.url = f"http://127.0.0.1:{self.httpd.server_address[1]}"
        threading.Thread(target=self.httpd.serve_forever, daemon=True).start()

    def tearDown(self):
        self.httpd.shutdown()
        self.httpd.server_close()
        shutil.rmtree(self.tmp, ignore_errors=True)

    def get(self, ruta, json_=True):
        with urllib.request.urlopen(self.url + ruta) as r:
            cuerpo = r.read().decode("utf-8")
            return json.loads(cuerpo) if json_ else cuerpo

    def accion(self, accion, version=None):
        if version is None:
            version = self.get("/api/presentador")["version"]
        datos = json.dumps({"version": version, "accion": accion}).encode()
        req = urllib.request.Request(self.url + "/api/accion", data=datos, headers={"Content-Type": "application/json"})
        try:
            with urllib.request.urlopen(req) as r:
                return r.status, json.loads(r.read())
        except urllib.error.HTTPError as e:
            with e:
                return e.code, json.loads(e.read())

    def test_paginas_disponibles(self):
        for ruta in ("/", "/publico", "/presentador", "/css/estilos.css", "/js/publico.js", "/js/presentador.js", "/js/sonidos.js",
                     "/js/espectaculos.js"):
            self.assertTrue(len(self.get(ruta, False)) > 100, ruta)
        with self.assertRaises(urllib.error.HTTPError) as ctx:
            self.get("/../datos/preguntas_procesadas.json", False)
        ctx.exception.close()

    def test_pagina_publica_no_tiene_controles(self):
        html = self.get("/publico", False) + self.get("/js/publico.js", False)
        self.assertNotIn("/api/accion", html)
        self.assertNotIn("/api/presentador", html)
        self.assertNotIn("/api/preguntas", html)

    def preparar(self, nombres=("Ana", "Luis", "Caro", "Pedro")):
        for accion in ({"tipo": "agregar_participantes", "nombres": list(nombres)}, {"tipo": "sortear_equipos"},
                       {"tipo": "mezclar_preguntas"}, {"tipo": "confirmar_partida"}, {"tipo": "sortear_lideres"}):
            codigo, r = self.accion(accion)
            self.assertEqual(codigo, 200, r.get("error"))

    def test_flujo_completo_sincronizado(self):
        datos = self.get("/api/preguntas")
        p = next(x for x in datos["preguntas"] if x["numero"] == 4)
        ids = [a["id"] for a in p["respuestas"]]
        self.assertEqual(p["respuestas"][0]["frecuencia"], 28)  # el panel recibe frecuencias y puntos
        self.preparar()
        self.assertEqual(self.accion({"tipo": "iniciar_ronda", "pregunta_id": p["id"]})[0], 200)
        pub = json.dumps(self.get("/api/publico"), ensure_ascii=False)
        for a in p["respuestas"]:
            self.assertNotIn(a["etiqueta"], pub)        # nada oculto en la vista pública
        self.accion({"tipo": "primero_en_responder", "equipo": 1})
        self.accion({"tipo": "respuesta_cara_a_cara", "equipo": 1, "respuesta_id": ids[0]})
        pub = self.get("/api/publico")
        pts = [a["puntos"] for a in p["respuestas"]]
        self.assertEqual(pub["ronda"]["casillas"][0], {"posicion": 1, "revelada": True, "texto": "Había trancón",
                                                       "puntos": pts[0], "equipo": 1, "como": "cara_a_cara"})
        self.assertEqual(pub["ronda"]["banco"], pts[0])
        self.assertEqual(pub["fase"], "decision")      # la respuesta n.º 1 gana el cara a cara automáticamente
        self.accion({"tipo": "decidir", "jugar": True})
        v = self.get("/api/presentador")["version"]   # doble clic: misma versión dos veces → la segunda se rechaza
        self.assertEqual(self.accion({"tipo": "strike"}, v)[0], 200)
        self.assertEqual(self.accion({"tipo": "strike"}, v)[0], 409)
        self.assertEqual(self.get("/api/publico")["ronda"]["strikes"], 1)
        self.accion({"tipo": "strike"})
        self.accion({"tipo": "strike"})
        self.assertEqual(self.get("/api/publico")["fase"], "robo")
        self.accion({"tipo": "resolver_robo", "respuesta_id": ids[1]})  # robo exitoso: n.º 1 + n.º 2 (una sola vez)
        self.accion({"tipo": "asignar_banco"})
        self.assertEqual(self.get("/api/publico")["puntajes"], [pts[0] + pts[1], 0])
        self.assertEqual(self.accion({"tipo": "asignar_banco"})[0], 409)
        self.assertEqual(self.get("/api/publico")["puntajes"], [pts[0] + pts[1], 0])

    def test_16_recargar_o_reiniciar_conserva_toda_la_partida(self):
        self.preparar(["Ana", "Luis", "Caro", "Pedro", "Sofía", "Juan"])
        self.accion({"tipo": "iniciar_ronda"})
        self.accion({"tipo": "primero_en_responder", "equipo": 0})
        pid = self.get("/api/presentador")["ronda"]["pregunta_id"]
        primera = next(x for x in self.get("/api/preguntas")["preguntas"] if x["id"] == pid)["respuestas"][0]["id"]
        self.accion({"tipo": "respuesta_cara_a_cara", "equipo": 0, "respuesta_id": primera})
        self.accion({"tipo": "ganador_cara_a_cara", "equipo": 0})
        self.accion({"tipo": "decidir", "jugar": True})
        self.accion({"tipo": "strike"})
        antes = self.get("/api/presentador")
        # Recargar la página = volver a pedir el estado; reiniciar el servidor = leer el archivo guardado.
        self.assertEqual(self.get("/api/presentador")["version"], antes["version"])
        otro = servidor.Servicio(JSON_PROCESADO, self.partida)
        despues = otro.presentador()
        for clave in ("preparacion", "lideres", "ronda", "puntajes", "preguntas_usadas", "fase", "historial_rondas"):
            self.assertEqual(despues[clave], antes[clave], clave)
        self.assertEqual(despues["config"]["equipos"], antes["config"]["equipos"])
        self.assertEqual(despues["ronda"]["strikes"], 1)
        self.assertIn(primera, despues["ronda"]["reveladas"])
        self.assertTrue(otro.presentador()["puede_deshacer"])

    def test_pantalla_publica_no_recibe_el_orden_de_las_preguntas(self):
        self.preparar()
        texto = json.dumps(self.get("/api/publico"), ensure_ascii=False)
        orden = self.get("/api/presentador")["preparacion"]["orden_preguntas"]
        self.assertTrue(orden)
        for pid in orden:
            self.assertNotIn(f'"{pid}"', texto)

    def test_error_claro_si_accion_no_valida(self):
        codigo, r = self.accion({"tipo": "strike"})
        self.assertEqual(codigo, 409)
        self.assertIn("Fase actual", r["error"])

    def test_espera_larga_entrega_el_cambio_al_instante(self):
        """La pantalla pública queda esperando y recibe la acción del presentador en cuanto ocurre."""
        import time
        v = self.get("/api/publico")["version"]
        resultado = {}

        def esperar():
            inicio = time.monotonic()
            resultado["estado"] = self.get(f"/api/publico?esperar={v}")
            resultado["segundos"] = time.monotonic() - inicio

        hilo = threading.Thread(target=esperar)
        hilo.start()
        time.sleep(0.4)
        self.assertTrue(hilo.is_alive())  # sigue esperando: no hubo cambios
        self.accion({"tipo": "configurar", "cambios": {"equipos": [{"nombre": "Rayos"}, {"nombre": "Truenos"}]}})
        hilo.join(5)
        self.assertFalse(hilo.is_alive())
        self.assertEqual(resultado["estado"]["equipos"][0]["nombre"], "Rayos")
        self.assertLess(resultado["segundos"], 3)

    def test_espera_larga_sin_cambios_devuelve_el_mismo_estado(self):
        v = self.servicio.publico()["version"]
        self.assertEqual(self.servicio.publico(v, espera=0.2)["version"], v)

    def test_controles_de_audio_llegan_a_la_pantalla_publica(self):
        self.accion({"tipo": "configurar", "cambios": {"volumen": 40, "musica": False}})
        pub = self.get("/api/publico")
        self.assertEqual((pub["volumen"], pub["musica"], pub["efectos"], pub["sonido"]), (40, False, True, True))
        self.accion({"tipo": "probar_sonido", "nombre": "correcta"})
        self.assertEqual(self.get("/api/publico")["evento"]["extra"]["sonido"], "correcta")


if __name__ == "__main__":
    unittest.main()
